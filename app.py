# -*- coding: utf-8 -*-
"""Công cụ phòng MT: đọc file PO (đơn đặt hàng siêu thị) → điền Phiếu Xuất Kho (PXK) của Nàng Cơm + tải trọng.

Web app cục bộ (127.0.0.1). Chạy: nhấp đúp Mo-cong-cu.bat (hoặc: python app.py). Lần đầu cần mã kích hoạt 20 ngày
do quản trị viên (ADMIN) cấp. Dữ liệu (danh mục, khách, PO, PXK) nằm trong thư mục công cụ, không gửi ra ngoài.
"""
import json
import os
import re
import socket
import sys
import threading
import webbrowser
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
import urllib.request

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / 'thu-vien'))

import licensing  # noqa: E402  (cấp phép offline 20 ngày — bản sao trong thu-vien, xem README)

DATA_DIR = ROOT / 'Dữ liệu'
OUT_DIR = ROOT / 'Xuất file'
WEB_DIR = ROOT / 'giao-dien'
CATALOG_FILE = DATA_DIR / 'DANH MỤC HH KÊNH MT.xlsx'   # danh mục sản phẩm (barcode → tải trọng/quy cách/giá)
CUSTOMER_FILE = DATA_DIR / 'khach-hang.xlsx'            # danh mục khách/cửa hàng (cửa hàng → MST/địa chỉ)
PXK_TEMPLATE = DATA_DIR / 'PXK-mau.xlsx'               # mẫu phiếu xuất kho

PORT = 8767
MAX_BODY = 20 * 1024 * 1024  # PO/PXK có thể lớn
APP_ID = 'mt'
LICENSE = licensing.LicenseClient()
LOCK = threading.RLock()


# =====================================================================
# danh mục sản phẩm (barcode → tải trọng/quy cách/giá) + khách/cửa hàng
# =====================================================================
from openpyxl import Workbook, load_workbook  # noqa: E402


def only_digits(s):
    return re.sub(r'\D', '', str(s or ''))


def _norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def _key(s):
    return _norm(s).lower()


def _header_cols(ws, needed, max_scan=4):
    """Dò dòng tiêu đề + cột theo tên (khớp chứa chuỗi, ưu tiên theo thứ tự). needed: {field: [chuỗi con]}."""
    for r in range(1, max_scan + 1):
        cells = {c: _key(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
        pos = {}
        for field, subs in needed.items():
            for sub in subs:
                col = next((c for c, t in cells.items() if t and sub in t and c not in pos.values()), None)
                if col:
                    pos[field] = col
                    break
        if all(f in pos for f in needed):
            return r, pos
    raise RuntimeError('File danh mục thiếu cột tiêu đề cần thiết.')


CAT_COLS = {'barcode': ['barcode', 'mã vạch'], 'ten': ['tên hàng', 'tên sản phẩm'], 'dvt': ['đvt'],
            'quy_cach': ['quy cách'], 'gia_chai': ['đơn giá chai', 'giá chai'], 'gia_thung': ['đơn giá thùng'],
            'tai_trong': ['tải trọng']}


def _num(v):
    if v is None or v == '':
        return 0.0
    try:
        return float(str(v).replace(',', ''))
    except (ValueError, TypeError):
        return 0.0


def load_catalog(path):
    """Đọc danh mục sản phẩm → {barcode: {barcode,ten,dvt,quy_cach,gia_chai,gia_thung,tai_trong}}."""
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hr, pos = _header_cols(ws, CAT_COLS)
    out = {}
    for r in range(hr + 1, ws.max_row + 1):
        bc = only_digits(ws.cell(r, pos['barcode']).value)
        if not bc:
            continue
        out[bc] = {'barcode': bc, 'ten': _norm(ws.cell(r, pos['ten']).value),
                   'dvt': _norm(ws.cell(r, pos['dvt']).value), 'quy_cach': int(_num(ws.cell(r, pos['quy_cach']).value)) or 1,
                   'gia_chai': _num(ws.cell(r, pos['gia_chai']).value), 'gia_thung': _num(ws.cell(r, pos['gia_thung']).value),
                   'tai_trong': _num(ws.cell(r, pos['tai_trong']).value)}
    wb.close()
    return out


def _backup(path):
    path = Path(path)
    if path.exists():
        bak = path.with_name(f'{path.stem}-saoluu-{datetime.now():%Y%m%d-%H%M%S}{path.suffix}')
        try:
            bak.write_bytes(path.read_bytes())
        except OSError:
            pass


def save_catalog(path, items):
    _backup(path)
    wb = Workbook()
    ws = wb.active
    ws.title = 'TẢI TRỌNG'
    ws.append(['STT', 'BARCODE', 'Tên Hàng', 'ĐVT', 'Quy Cách', 'Đơn Giá Chai', 'Đơn Giá Thùng', 'Tải trọng'])
    for i, p in enumerate(items.values(), 1):
        ws.append([i, p['barcode'], p['ten'], p['dvt'], p['quy_cach'], p['gia_chai'], p['gia_thung'], p['tai_trong']])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def add_item(items, d):
    bc = only_digits(d['barcode'])
    items[bc] = {'barcode': bc, 'ten': _norm(d.get('ten')), 'dvt': _norm(d.get('dvt')),
                 'quy_cach': int(_num(d.get('quy_cach'))) or 1, 'gia_chai': _num(d.get('gia_chai')),
                 'gia_thung': _num(d.get('gia_thung')), 'tai_trong': _num(d.get('tai_trong'))}
    return items[bc]


def update_item(items, barcode, d):
    p = items[only_digits(barcode)]
    for k in ('ten', 'dvt'):
        if k in d:
            p[k] = _norm(d[k])
    for k in ('quy_cach',):
        if k in d:
            p[k] = int(_num(d[k])) or p[k]
    for k in ('gia_chai', 'gia_thung', 'tai_trong'):
        if k in d:
            p[k] = _num(d[k])
    return p


def delete_item(items, barcode):
    items.pop(only_digits(barcode), None)


# ---- khách / cửa hàng ----
CUST_COLS = {'store': ['cửa hàng', 'store'], 'ten_phap_nhan': ['tên pháp nhân', 'đơn vị nhận', 'pháp nhân'],
             'mst': ['mst', 'mã số thuế'], 'dia_chi': ['địa chỉ']}


def load_customers(path):
    """Đọc danh mục khách/cửa hàng → {key cửa hàng: {store,ten_phap_nhan,mst,dia_chi}}. Chưa có file → {}."""
    if not Path(path).exists():
        return {}
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    try:
        hr, pos = _header_cols(ws, CUST_COLS)
    except RuntimeError:
        wb.close()
        return {}
    out = {}
    for r in range(hr + 1, ws.max_row + 1):
        store = _norm(ws.cell(r, pos['store']).value)
        if not store:
            continue
        out[_key(store)] = {'store': store, 'ten_phap_nhan': _norm(ws.cell(r, pos['ten_phap_nhan']).value),
                            'mst': only_digits(ws.cell(r, pos['mst']).value), 'dia_chi': _norm(ws.cell(r, pos['dia_chi']).value)}
    wb.close()
    return out


def save_customers(path, customers):
    _backup(path)
    wb = Workbook()
    ws = wb.active
    ws.title = 'KHÁCH'
    ws.append(['Cửa hàng', 'Tên pháp nhân', 'MST', 'Địa chỉ'])
    for c in customers.values():
        ws.append([c['store'], c['ten_phap_nhan'], c['mst'], c['dia_chi']])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def add_customer(customers, d):
    store = _norm(d['store'])
    customers[_key(store)] = {'store': store, 'ten_phap_nhan': _norm(d.get('ten_phap_nhan')),
                              'mst': only_digits(d.get('mst')), 'dia_chi': _norm(d.get('dia_chi'))}
    return customers[_key(store)]


def _deaccent(s):
    import unicodedata
    s = unicodedata.normalize('NFD', _key(s)).replace('đ', 'd')
    return ''.join(c for c in s if unicodedata.category(c) != 'Mn')


def _words(s):
    return set(_deaccent(s).split())


def match_customer(customers, store_hint):
    """Khớp cửa hàng. Chỉ nhận khi CHẮC CHẮN: trùng đúng, hoặc một bên là tập con từ của bên kia VÀ duy nhất
    (so sánh bỏ dấu). Mơ hồ (2 chi nhánh cùng tiền tố, hoặc chỉ trùng vài từ) -> None để người dùng tự điền
    (tránh gán nhầm MST/địa chỉ)."""
    k = _key(store_hint)
    if k in customers:
        return customers[k]
    want = _words(store_hint)
    if not want:
        return None
    cands = [c for key, c in customers.items() if (want <= _words(key)) or (_words(key) <= want)]
    return cands[0] if len(cands) == 1 else None


# =====================================================================
# lõi: đơn vị → tổng đơn vị lẻ → số thùng → ghép danh mục → tải trọng
# =====================================================================
VAT = 1.08  # đơn giá danh mục đã gồm VAT; cột đơn giá PXK là chưa VAT -> chia 1.08


def to_le(so_luong, dvt_po, per_box, unit_kind, quy_cach=None):
    """Tổng số đơn vị LẺ (chai/lọ/hũ...). box/thùng → số lượng × (per_box ghi trong PO, hoặc quy_cach danh mục);
    lẻ → số lượng."""
    so_luong = _num(so_luong)
    if unit_kind == 'box':
        per = per_box or quy_cach or 1
        return so_luong * per
    return so_luong


def match_line(line, catalog):
    """Ghép 1 dòng PO với danh mục theo barcode. Trả dict kết quả; mã lạ -> matched False."""
    bc = only_digits(line.get('barcode'))
    p = catalog.get(bc)
    if not p and len(bc) == 13 and bc[:12] in catalog:  # mã 12 số trong PDF dính thêm 1 số -> thử 12 số
        bc, p = bc[:12], catalog[bc[:12]]
    if not p:
        return {'barcode': bc, 'matched': False, 'so_luong': _num(line.get('so_luong')),
                'dvt_po': line.get('dvt_po'), 'tai_trong_dong': 0, 'so_thung': 0, 'tong_le': 0}
    quy_cach = p['quy_cach'] or 1
    tong_le = to_le(line.get('so_luong'), line.get('dvt_po'), line.get('per_box'), line.get('unit_kind', 'le'), quy_cach)
    so_thung = tong_le / quy_cach if quy_cach else tong_le
    don_gia = p['gia_thung'] / VAT  # đơn giá thùng chưa VAT
    return {'barcode': bc, 'matched': True, 'ten': p['ten'], 'dvt': p['dvt'], 'quy_cach': quy_cach,
            'tong_le': tong_le, 'so_thung': so_thung, 'tai_trong_dong': so_thung * p['tai_trong'],
            'don_gia': don_gia, 'thanh_tien': so_thung * don_gia, 'so_luong': _num(line.get('so_luong')),
            'dvt_po': line.get('dvt_po')}


def build_lines(lines, catalog):
    """Ghép mọi dòng; trả (rows, unknown=list barcode chưa có trong danh mục)."""
    rows = [match_line(ln, catalog) for ln in lines]
    unknown = [r['barcode'] for r in rows if not r['matched'] and r['barcode']]
    return rows, unknown


def tong_tai_trong(rows):
    return sum(r['tai_trong_dong'] for r in rows if r.get('matched'))


# =====================================================================
# điền Phiếu Xuất Kho (bản sao mẫu; dựng lại khối TỔNG/VAT/chữ ký theo số dòng)
# =====================================================================
# Vị trí cố định theo file mẫu "PHIEU XUAT KHO": tiêu đề dòng 13, dòng hàng đầu 14, mẫu có 4 dòng (14–17),
# TỔNG dòng 18, VAT 19–21, khối chữ ký 24–29.
PXK_SHEET = 'PHIEU XUAT KHO'
_HDR, _FIRST, _ORIG_N = 13, 14, 4
_ORIG_TONG, _ORIG_VAT, _SIG = 18, (19, 20, 21), (24, 29)


def _sanitize(s):
    return re.sub(r'[\\/:*?"<>|]+', '_', str(s or '')).strip() or 'PO'


def fill_pxk(parsed, catalog, template_path, out_dir, kho='KHO-BENTRE', ngay=None):
    """Điền ParsedPO vào bản sao mẫu PXK. Trả (out_path, unknown). Chỉ dòng khớp danh mục lên PXK."""
    from copy import copy
    template_path = Path(template_path)
    if not template_path.exists():
        raise RuntimeError(f'Không tìm thấy file mẫu PXK: {template_path}')
    rows_all, unknown = build_lines(parsed.get('lines', []), catalog)
    rows = [r for r in rows_all if r.get('matched')]
    if not rows:
        raise RuntimeError('Không có dòng hàng nào khớp danh mục để tạo PXK.')
    ngay = ngay or datetime.now()
    h = parsed.get('header', {})
    so_po = h.get('so_po') or ''
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    out = Path(out_dir) / f'PXK {_sanitize(so_po)} {ngay:%Y%m%d}.xlsx'
    import shutil as _sh
    _sh.copyfile(template_path, out)
    wb = load_workbook(out)
    if PXK_SHEET not in wb.sheetnames:
        wb.close()
        raise RuntimeError(f'File mẫu thiếu sheet "{PXK_SHEET}".')
    ws = wb[PXK_SHEET]
    N = len(rows)
    delta = N - _ORIG_N

    data_style = {c: copy(ws.cell(_FIRST, c)._style) for c in range(2, 14)}
    tong_style = {c: copy(ws.cell(_ORIG_TONG, c)._style) for c in range(2, 14)}
    vat_style = {rw: {c: copy(ws.cell(rw, c)._style) for c in range(2, 14)} for rw in _ORIG_VAT}
    sig = [(rr, c, ws.cell(rr, c).value, copy(ws.cell(rr, c)._style))
           for rr in range(_SIG[0], _SIG[1] + 1) for c in range(2, 14)]
    orig_merges = [(m.min_col, m.min_row, m.max_col, m.max_row) for m in ws.merged_cells.ranges]

    for m in list(ws.merged_cells.ranges):   # bỏ mọi ô gộp từ dòng 14 trở xuống (sẽ dựng lại theo vị trí mới)
        if m.min_row >= _FIRST:
            ws.unmerge_cells(str(m))
    for rr in range(_FIRST, ws.max_row + 1):
        for c in range(2, 14):
            ws.cell(rr, c).value = None

    for i, r in enumerate(rows):            # dòng hàng
        rr = _FIRST + i
        vals = {2: i + 1, 3: kho, 5: r['barcode'], 6: r['ten'], 7: r['dvt'], 8: r['quy_cach'],
                9: r['so_thung'], 10: f'=H{rr}*I{rr}', 11: round(r['don_gia'], 2), 12: f'=K{rr}*I{rr}'}
        for c in range(2, 14):
            cell = ws.cell(rr, c)
            cell.value = vals.get(c)
            cell._style = copy(data_style[c])
    last = _FIRST + N - 1
    tong = _FIRST + N
    for c in range(2, 14):
        ws.cell(tong, c)._style = copy(tong_style[c])
    ws.cell(tong, 2).value = 'TỔNG:'
    ws.cell(tong, 9).value = f'=SUM(I14:I{last})'
    ws.cell(tong, 10).value = f'=SUM(J14:J{last})'
    vat1, vat2, vat3 = tong + 1, tong + 2, tong + 3
    for new, orig in zip((vat1, vat2, vat3), _ORIG_VAT):
        for c in range(2, 14):
            ws.cell(new, c)._style = copy(vat_style[orig][c])
    ws.cell(vat1, 10).value = 'TỔNG CỘNG CHƯA BAO GỒM VAT:'
    ws.cell(vat1, 13).value = f'=SUM(L14:L{last})'
    ws.cell(vat2, 10).value = 'THUẾ VAT:'
    ws.cell(vat2, 12).value = 0.08
    ws.cell(vat2, 13).value = f'=M{vat1}*L{vat2}'
    ws.cell(vat3, 10).value = 'TỔNG GIÁ TRỊ ĐƠN HÀNG ĐÃ BAO GỒM VAT:'
    ws.cell(vat3, 13).value = f'=M{vat1}+M{vat2}'
    for rr, c, val, style in sig:           # khối chữ ký dời theo delta
        cell = ws.cell(rr + delta, c)
        cell.value = val
        cell._style = style
    from openpyxl.utils import get_column_letter
    for minc, minr, maxc, maxr in orig_merges:   # dựng lại ô gộp (dòng ≥14 dời theo delta)
        if minr >= _FIRST:
            ws.merge_cells(f'{get_column_letter(minc)}{minr + delta}:{get_column_letter(maxc)}{maxr + delta}')

    # header
    don_vi = h.get('don_vi_nhan') or ''
    if h.get('mst'):
        don_vi = f'{don_vi} - MST:{h["mst"]}'
    ws.cell(5, 2).value = f'Ngày {ngay:%d} tháng {ngay:%m} năm {ngay:%Y}'
    ws.cell(7, 2).value = f'Đơn vị nhận hàng: {don_vi}'
    ws.cell(8, 2).value = f'Người nhận hàng: {h.get("nguoi_nhan") or ""}'
    ws.cell(9, 2).value = f'Địa chỉ nhận hàng: {h.get("dia_chi") or ""}'
    ws.cell(10, 2).value = f'Lý do xuất: {h.get("ly_do") or "Xuất kho bán hàng"}'
    ws.cell(9, 10).value = f'Đơn hàng(PO) số: {so_po}'
    ws.cell(1, 10).value = f'{so_po}\nSố: PXK/{kho}-{ngay:%d%m%y}'  # số phiếu theo PO này (thay số mẫu EMART)
    w = tong_tai_trong(rows)
    ws.cell(11, 12).value = f'Khối lượng:  {w:g} Kg'
    wb.save(out)
    wb.close()
    return str(out), unknown


# =====================================================================
# đọc PO: PDF (pypdf) + nhận diện chuỗi + parser từng chuỗi
# =====================================================================
# Mã vạch Nàng Cơm: 893 + 9–10 số (12–13 số). Trích theo prefix nên đọc đúng cả khi dính liền mã hàng nội bộ.
_BARCODE = re.compile(r'893\d{9,10}')


def find_barcodes(text):
    return _BARCODE.findall(str(text or ''))


def pdf_text(path):
    import pypdf
    r = pypdf.PdfReader(str(path))
    return '\n'.join((p.extract_text() or '') for p in r.pages)


def detect_chain(path):
    """Nhận diện chuỗi theo đuôi file + dấu hiệu nội dung."""
    path = Path(path)
    if path.suffix.lower() in ('.xlsx', '.xlsm', '.xls'):
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            names = ' '.join(wb.sheetnames).lower()
            head = ' '.join(str(c.value) for row in wb.worksheets[0].iter_rows(max_row=12) for c in row if c.value).lower()
            wb.close()
        except Exception:
            names = head = ''
        if 'timedelivery' in names or 'kingkong' in head:
            return 'kingkong'
        return 'excel_unknown'
    text = pdf_text(path).lower()
    for key, chain in (('thiso', 'emart'), ('emart', 'emart'), ('menas', 'menas'), ('mgm', 'menas'),
                       ('fujimart', 'fujimart'), ('maxidi', 'bd'), ('bưu điện', 'bhbd'), ('vietnam post', 'bhbd')):
        if key in text:
            return chain
    return 'unknown'


def _emart_lines(text):
    pos = [(m.start(), m.group()) for m in _BARCODE.finditer(text)]
    ends = [p[0] for p in pos[1:]] + [len(text)]
    lines = []
    for (start, bc), end in zip(pos, ends):
        win = text[start:end]
        m = re.search(r'BOX\s+(\d+)\s+(\d+)', win)
        per_box = int(m.group(1)) if m else None
        so_luong = int(m.group(2)) if m else 0
        lines.append({'barcode': bc, 'so_luong': so_luong, 'dvt_po': 'BOX', 'unit_kind': 'box', 'per_box': per_box})
    return lines


def parse_emart(text):
    so_po = re.search(r'PO No\.?\s*:?\s*(\d+)', text)
    deliv = re.search(r'Delivery to\s*:?\s*(.+?)\s+Delivery Date', text, re.S)
    don_vi = re.search(r'(.*(?:RETAIL|COMPANY LIMITED).*)', text)
    nguoi = _norm(deliv.group(1)) if deliv else ''
    store = re.match(r'(EMART\s+\S+)', nguoi)
    return {'chain': 'emart', 'lines': _emart_lines(text),
            'header': {'so_po': so_po.group(1) if so_po else '', 'nguoi_nhan': nguoi, 'store': store.group(1) if store else nguoi,
                       'don_vi_nhan': _norm(don_vi.group(1)) if don_vi else '', 'mst': '', 'dia_chi': nguoi,
                       'ngay_dat': '', 'ly_do': 'Xuất kho bán hàng'}}


def parse_menas(text):
    """MENAS/MGM (PDF): mỗi mặt hàng là 1 khối dọc: ...barcode, tên, ĐVT (GOI/HOP/LON), đơn giá, SỐ LƯỢNG, ..."""
    L = [l.strip() for l in text.split('\n') if l.strip()]
    lines = []
    for i, l in enumerate(L):
        if not re.fullmatch(r'893\d{9,10}', l):
            continue
        j = i + 1
        while j < len(L) and not re.fullmatch(r'[A-Z]{2,4}', L[j]):  # bỏ qua dòng tên (chữ thường/hỗn hợp)
            j += 1
        if j + 2 >= len(L):
            continue
        dvt = L[j]
        so_luong = int(re.sub(r'\D', '', L[j + 2]) or 0)  # sau ĐVT là đơn giá rồi SỐ LƯỢNG
        lines.append({'barcode': l, 'so_luong': so_luong, 'dvt_po': dvt, 'unit_kind': 'le'})
    so_po = re.search(r'PO\d{3,}-\d{3,}', text)
    ship = re.search(r'MGM\s+\S+', text)
    return {'chain': 'menas', 'lines': lines,
            'header': {'so_po': so_po.group(0) if so_po else '', 'nguoi_nhan': ship.group(0) if ship else '',
                       'store': ship.group(0) if ship else '', 'don_vi_nhan': 'CÔNG TY TNHH MENAS', 'mst': '',
                       'dia_chi': '', 'ngay_dat': '', 'ly_do': 'Xuất kho bán hàng'}}


def parse_fujimart(text):
    """FujiMart (PDF TCVN3): mỗi dòng 'STT qty.0 tiền ĐVT đơn_giá tên ...barcode(dính mã hàng)'."""
    lines = []
    for l in text.split('\n'):
        bcs = find_barcodes(l)
        m = re.match(r'\s*\d+\s+(\d+)(?:\.\d+)?\s', l)  # số lượng = số thứ 2 đầu dòng
        if bcs and m:
            dvt = re.search(r'\b(LO|CHA|CAN|LON|BINH)\b', l)
            lines.append({'barcode': bcs[0], 'so_luong': int(m.group(1)), 'dvt_po': dvt.group(1) if dvt else 'LO',
                          'unit_kind': 'le'})
    so_don = re.search(r'\b(13\d{16,})\b', text)  # số đơn dạng 1320033026...
    store = 'FujiMart' + (' Linh Đàm' if 'linh' in text.lower() or 'Linh' in text else '')
    return {'chain': 'fujimart', 'lines': lines,
            'header': {'so_po': so_don.group(1) if so_don else '', 'nguoi_nhan': store, 'store': store,
                       'don_vi_nhan': 'CÔNG TY TNHH BÁN LẺ FUJIMART VIỆT NAM', 'mst': '', 'dia_chi': '',
                       'ngay_dat': '', 'ly_do': 'Xuất kho bán hàng'}}


def parse_bhbd(text):
    """BHBĐ/Nam Trực (PDF): dòng '... barcode qty ĐVT ...'."""
    lines = []
    for l in text.split('\n'):
        m = re.search(r'(893\d{9,10})\s+(\d+)\s+([A-ZÀ-Ỹ]{2,6})', l)
        if m:
            lines.append({'barcode': m.group(1), 'so_luong': int(m.group(2)), 'dvt_po': m.group(3), 'unit_kind': 'le'})
    so_po = re.search(r'PO No\.?\)?\s*:?\s*(\S+)', text) or re.search(r'(PDN\d+)', text)
    return {'chain': 'bhbd', 'lines': lines,
            'header': {'so_po': so_po.group(1) if so_po else '', 'nguoi_nhan': 'Bưu điện Việt Nam', 'store': 'Nam Trực',
                       'don_vi_nhan': '', 'mst': '', 'dia_chi': '', 'ngay_dat': '', 'ly_do': 'Xuất kho bán hàng'}}


def parse_kingkong(path):
    """KingKong (Excel sheet 'timedelivery'): cột D mã vạch, K số lượng, L ĐVT; header D6/F11/F12."""
    wb = load_workbook(path, data_only=True)
    ws = wb['timedelivery'] if 'timedelivery' in wb.sheetnames else wb.worksheets[0]
    lines = []
    for r in range(1, ws.max_row + 1):
        bc = only_digits(ws.cell(r, 4).value)
        if re.fullmatch(r'893\d{9,10}', bc or ''):
            lines.append({'barcode': bc, 'so_luong': int(_num(ws.cell(r, 11).value)), 'dvt_po': _norm(ws.cell(r, 12).value),
                          'unit_kind': 'le'})
    ma_phieu = re.search(r'(OS\d+)', str(ws['D6'].value or ''))
    store = _norm(ws['F12'].value) or _norm(ws['F11'].value)
    wb.close()
    return {'chain': 'kingkong', 'lines': lines,
            'header': {'so_po': ma_phieu.group(1) if ma_phieu else '', 'nguoi_nhan': store, 'store': store,
                       'don_vi_nhan': store, 'mst': '', 'dia_chi': '', 'ngay_dat': '', 'ly_do': 'Xuất kho bán hàng'}}


_PARSERS = {'emart': parse_emart, 'menas': parse_menas, 'fujimart': parse_fujimart, 'bhbd': parse_bhbd,
            'kingkong': parse_kingkong}


def parse_po(path):
    """Đọc 1 file PO → ParsedPO {header, lines, chain}. Chuỗi chưa hỗ trợ → RuntimeError."""
    chain = detect_chain(path)
    parser = _PARSERS.get(chain)
    if not parser:
        raise RuntimeError(f'Chưa hỗ trợ mẫu PO này ({chain}). Hãy nhập thông tin tay hoặc báo để bổ sung.')
    if Path(path).suffix.lower() in ('.xlsx', '.xlsm', '.xls'):
        return parser(path)
    return parser(pdf_text(path))


def apply_customer(parsed, customers):
    """Điền tên pháp nhân/MST/địa chỉ từ danh mục khách nếu tìm được theo cửa hàng."""
    h = parsed.get('header', {})
    c = match_customer(customers, h.get('store') or h.get('nguoi_nhan') or '')
    if c:
        h.setdefault('don_vi_nhan', '')
        if c.get('ten_phap_nhan'):
            h['don_vi_nhan'] = c['ten_phap_nhan']
        if c.get('mst'):
            h['mst'] = c['mst']
        if c.get('dia_chi'):
            h['dia_chi'] = c['dia_chi']
    return parsed


# =====================================================================
# endpoints
# =====================================================================

def _load_masters():
    cat = load_catalog(CATALOG_FILE) if CATALOG_FILE.exists() else {}
    cus = load_customers(CUSTOMER_FILE)
    return cat, cus


def handle(path, body):
    """Điều phối POST nghiệp vụ."""
    if path == '/api/po/parse':
        cat, cus = _load_masters()
        parsed = apply_customer(parse_po(Path(body['path'])), cus)
        rows, unknown = build_lines(parsed['lines'], cat)
        return {'header': parsed['header'], 'chain': parsed['chain'], 'rows': rows, 'unknown': unknown,
                'tai_trong': tong_tai_trong(rows)}
    if path == '/api/po/pxk':
        cat, cus = _load_masters()
        parsed = apply_customer(parse_po(Path(body['path'])), cus)
        if body.get('header'):
            parsed['header'].update(body['header'])  # cho sửa header ở preview
        if not PXK_TEMPLATE.exists():
            raise RuntimeError('Chưa có file mẫu PXK trong thư mục "Dữ liệu" (PXK-mau.xlsx).')
        out, unknown = fill_pxk(parsed, cat, PXK_TEMPLATE, OUT_DIR)
        return {'file': out, 'unknown': unknown}
    if path == '/api/po/list':
        files = []
        for d in (DATA_DIR / 'PO', DATA_DIR, ROOT):
            if d.exists():
                files += [str(p) for p in d.glob('*') if p.suffix.lower() in ('.pdf', '.xlsx', '.xls', '.xlsm')
                          and 'DANH MỤC' not in p.name and 'PXK' not in p.name.upper() and 'khach' not in p.name]
        return {'files': sorted(set(files))}
    if path == '/api/open':
        target = Path(body['path']).resolve()
        if target.exists() and (OUT_DIR.resolve() in target.parents or ROOT.resolve() in target.parents):
            os.startfile(target)  # noqa: S606
        return {'ok': True}
    # ---- CRUD danh mục sản phẩm ----
    if path == '/api/catalog/list':
        return {'items': list(_load_masters()[0].values())}
    if path in ('/api/catalog/add', '/api/catalog/update'):
        cat = _load_masters()[0]
        it = body['item']
        if path.endswith('update') and only_digits(it.get('barcode')) in cat:
            update_item(cat, it['barcode'], it)
        else:
            add_item(cat, it)
        save_catalog(CATALOG_FILE, cat)
        return {'ok': True}
    if path == '/api/catalog/delete':
        cat = _load_masters()[0]
        delete_item(cat, body['barcode'])
        save_catalog(CATALOG_FILE, cat)
        return {'ok': True}
    # ---- CRUD khách / cửa hàng ----
    if path == '/api/customers/list':
        return {'items': list(load_customers(CUSTOMER_FILE).values())}
    if path in ('/api/customers/add', '/api/customers/update'):
        cus = load_customers(CUSTOMER_FILE)
        add_customer(cus, body['item'])
        save_customers(CUSTOMER_FILE, cus)
        return {'ok': True}
    if path == '/api/customers/delete':
        cus = load_customers(CUSTOMER_FILE)
        cus.pop(_key(body['store']), None)
        save_customers(CUSTOMER_FILE, cus)
        return {'ok': True}
    raise ValueError('Không rõ yêu cầu: ' + path)


# =====================================================================
# máy chủ web (gia cố như các công cụ nang-com khác)
# =====================================================================

def license_info():
    return dict(LICENSE.check_license().to_dict(), policy=licensing.POLICY_TEXT, app=APP_ID)


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _send(self, code, payload, ctype='application/json; charset=utf-8'):
        data = payload if isinstance(payload, bytes) else json.dumps(payload, ensure_ascii=False, default=str).encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', ctype)
        self.send_header('Cache-Control', 'no-store')
        self.end_headers()
        self.wfile.write(data)

    def _host_ok(self):
        host = (self.headers.get('Host') or '').strip().lower()
        host = host[1:host.index(']')] if host.startswith('[') and ']' in host else host.rsplit(':', 1)[0]
        if host in ('127.0.0.1', 'localhost', '::1'):
            return True
        self._send(403, {'error': 'Yêu cầu không hợp lệ.'})
        return False

    def _read_json(self, limit=MAX_BODY):
        if not (self.headers.get('Content-Type') or '').startswith('application/json'):
            raise ValueError('content-type')
        n = int(self.headers.get('Content-Length') or 0)
        if n < 0 or n > limit:
            raise ValueError('size')
        body = json.loads(self.rfile.read(n) or b'{}')
        if not isinstance(body, dict):
            raise ValueError('body')
        return body

    def _license_blocked(self):
        st = LICENSE.check_license()
        if st.valid:
            return False
        self._send(403, {'error': st.message, 'license': license_info()})
        return True

    def _activate(self):
        try:
            body = self._read_json(limit=8192)
        except ValueError as e:
            return self._send(415 if str(e) == 'content-type' else 400, {'error': 'Yêu cầu không hợp lệ.'})
        try:
            st = LICENSE.activate(str(body.get('code') or ''))
        except licensing.ActivationError as e:
            return self._send(400, {'error': str(e), 'reason': e.reason, 'license': license_info()})
        except (ValueError, AttributeError):
            return self._send(400, {'error': 'Yêu cầu không hợp lệ.'})
        print('[License]', st.message)
        self._send(200, {'ok': True, 'license': license_info()})

    def do_GET(self):
        if not self._host_ok():
            return
        u = urlparse(self.path)
        if u.path == '/api/license':
            return self._send(200, license_info())
        if u.path.startswith('/api/'):
            if self._license_blocked():
                return
            return self._send(404, {'error': 'Không tìm thấy'})
        name = 'index.html' if u.path in ('/', '') else u.path.lstrip('/')
        f = (WEB_DIR / name).resolve()
        if WEB_DIR.resolve() not in f.parents or not f.is_file():
            return self._send(404, {'error': 'Không tìm thấy'})
        ctype = {'.html': 'text/html', '.css': 'text/css', '.js': 'application/javascript'}.get(f.suffix, 'application/octet-stream')
        self._send(200, f.read_bytes(), ctype + '; charset=utf-8')

    def do_POST(self):
        if not self._host_ok():
            return
        u = urlparse(self.path)
        if u.path == '/api/license/activate':
            return self._activate()
        if self._license_blocked():
            return
        try:
            body = self._read_json()
        except ValueError as e:
            return self._send(415 if str(e) == 'content-type' else 400, {'error': 'Yêu cầu không hợp lệ.'})
        try:
            with LOCK:
                res = handle(u.path, body)
            self._send(200, res)
        except (ValueError, RuntimeError, KeyError, FileNotFoundError) as e:
            self._send(400, {'error': str(e) or 'Dữ liệu không hợp lệ.'})
        except Exception as e:
            self._send(500, {'error': f'Lỗi không mong muốn: {e}.'})


class Server(ThreadingHTTPServer):
    allow_reuse_address = False

    def server_bind(self):
        if hasattr(socket, 'SO_EXCLUSIVEADDRUSE'):
            self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        super().server_bind()


def already_running(port):
    try:
        with urllib.request.urlopen(f'http://127.0.0.1:{port}/api/license', timeout=2) as r:
            return json.loads(r.read()).get('app') == APP_ID
    except Exception:
        return False


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    print('Đang mở công cụ MT (PO → Phiếu Xuất Kho)...')
    lic = LICENSE.check_license()
    print('[License]', lic.message)
    if not lic.valid:
        print('Nhập mã kích hoạt trên giao diện. Mã thiết bị của máy này:', lic.device_id or '-')
    DATA_DIR.mkdir(exist_ok=True)
    OUT_DIR.mkdir(exist_ok=True)
    server = None
    for port in range(PORT, PORT + 20):
        try:
            server = Server(('127.0.0.1', port), Handler)
            break
        except OSError:
            if already_running(port):
                url = f'http://127.0.0.1:{port}/'
                print('Công cụ ĐANG MỞ ở một cửa sổ khác — mở lại giao diện đó:', url)
                if '--khong-mo' not in sys.argv:
                    webbrowser.open(url)
                return
            continue
    if server is None:
        print(f'(!) Không mở được cổng {PORT}–{PORT + 19}. Đóng bớt chương trình rồi mở lại.')
        input('Bấm Enter để đóng...')
        return
    url = f'http://127.0.0.1:{port}/'
    print('Giao diện:', url)
    print('ĐỪNG ĐÓNG cửa sổ này trong lúc dùng công cụ.')
    if '--khong-mo' not in sys.argv:
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    server.serve_forever()


if __name__ == '__main__':
    main()
